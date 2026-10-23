"""Artifact parts to model input, and the one limiter every prompt path uses.

``load_content_parts`` turns stored artifact parts into ``ContentPart`` values. ``fit_parts``
keeps a request inside the provider limits; it never truncates a part, it drops whole trailing
parts and says which ones, so a dropped part is always visible to the caller.
"""

import asyncio
import csv
import io
from collections.abc import Sequence
from dataclasses import dataclass
from typing import Any, Final, Literal, cast

from cashapp.integrations.llm.base import (
    ContentPart,
    ImagePart,
    LlmError,
    LlmErrorCode,
    PdfPart,
    TextPart,
)
from cashapp.integrations.storage.base import ObjectStore

_IMAGE_TYPES: Final = frozenset({"image/png", "image/jpeg", "image/gif", "image/webp"})
# Anthropic rejects requests over 32 MB including base64 expansion; stay under with margin.
MAX_ENCODED_BYTES: Final = 31_000_000
_PROVIDER: Final = "content"


@dataclass(frozen=True, slots=True)
class ArtifactPartRef:
    part_no: int
    kind: str
    media_type: str
    object_key: str
    purged: bool = False

    @classmethod
    def from_row(cls, row: Any) -> "ArtifactPartRef":
        """Duck-typed over the ``ArtifactPart`` ORM row (kind may be an enum)."""
        kind = row.kind
        return cls(
            part_no=int(row.part_no),
            kind=str(getattr(kind, "value", kind)),
            media_type=str(row.media_type),
            object_key=str(row.object_key),
            purged=row.purged_at is not None,
        )


def _fail(detail: str) -> LlmError:
    return LlmError(LlmErrorCode.EXTRACTION_FAILED, detail, provider=_PROVIDER, model_id="")


def _label(ref: ArtifactPartRef) -> str:
    return f"[part {ref.part_no}: {ref.kind.lower()}]\n"


def _docx_text(data: bytes) -> str:
    from docx import Document  # deferred: only DOCX parts need it

    document: Any = Document(io.BytesIO(data))
    lines: list[str] = [str(p.text) for p in document.paragraphs if str(p.text).strip()]
    for table in document.tables:
        for row in table.rows:
            lines.append("\t".join(str(cell.text).strip() for cell in row.cells))
    return "\n".join(lines)


def _as_supported_image(data: bytes) -> tuple[bytes, Literal["image/png"]]:
    """Re-encode an unsupported raster type (TIFF, BMP, ...) as PNG."""
    from PIL import Image

    with Image.open(io.BytesIO(data)) as image:
        image.seek(0)
        converted: Any = image.convert("RGB") if image.mode not in {"RGB", "RGBA", "L"} else image
        out = io.BytesIO()
        converted.save(out, format="PNG")
        return out.getvalue(), "image/png"


def _csv_text(raw: str) -> str:
    """Normalise a table part to CSV so cell boundaries survive."""
    try:
        rows = list(csv.reader(io.StringIO(raw)))
    except csv.Error:
        return raw
    out = io.StringIO()
    csv.writer(out, lineterminator="\n").writerows(rows)
    return out.getvalue()


async def _load_one(ref: ArtifactPartRef, store: ObjectStore) -> ContentPart:
    data = await store.get(ref.object_key)
    kind = ref.kind.upper()
    if kind in {"TEXT", "TABLE"}:
        decoded = data.decode("utf-8", errors="replace")
        body = _csv_text(decoded) if kind == "TABLE" else decoded
        return TextPart(_label(ref) + body)
    if kind == "DOCX":
        body = await asyncio.to_thread(_docx_text, data)
        return TextPart(_label(ref) + body)
    if kind == "PDF":
        return PdfPart(data)
    if kind == "IMAGE":
        if ref.media_type in _IMAGE_TYPES:
            return ImagePart(data, cast("Literal['image/png']", ref.media_type))
        try:
            png, media_type = await asyncio.to_thread(_as_supported_image, data)
        except Exception as exc:
            raise _fail(f"part {ref.part_no}: image cannot be decoded") from exc
        return ImagePart(png, media_type)
    raise _fail(f"part {ref.part_no}: unsupported part kind {ref.kind}")


async def load_content_parts(
    parts: Sequence[ArtifactPartRef], store: ObjectStore
) -> list[ContentPart]:
    """One ``ContentPart`` per artifact part, in the order given (callers pass part-number order).

    Raises ``LlmError(EXTRACTION_FAILED)`` for a purged or undecodable part.
    """
    for ref in parts:
        if ref.purged:
            raise _fail("artifact purged")
    return list(await asyncio.gather(*(_load_one(ref, store) for ref in parts)))


@dataclass(frozen=True, slots=True)
class FittedParts:
    parts: tuple[ContentPart, ...]
    dropped_part_nos: tuple[int, ...]


def _encoded_size(size: int) -> int:
    return 4 * ((size + 2) // 3)


def fit_parts(
    parts: Sequence[ContentPart],
    *,
    max_text_chars: int = 400_000,
    max_binary_bytes: int = 30_000_000,
    max_parts: int = 40,
    part_nos: Sequence[int] | None = None,
) -> FittedParts:
    """Keep leading parts while text, binary and count budgets hold; drop the rest whole.

    ``part_nos`` maps each position to its artifact part number for the dropped list; without it
    the 1-based position is reported.
    """
    text_chars = 0
    binary_bytes = 0
    encoded_bytes = 0
    kept: list[ContentPart] = []
    for part in parts:
        if isinstance(part, TextPart):
            text_chars += len(part.text)
            size = 0
        else:
            size = len(part.data)
            binary_bytes += size
            encoded_bytes += _encoded_size(size)
        if (
            len(kept) >= max_parts
            or text_chars > max_text_chars
            or binary_bytes > max_binary_bytes
            or encoded_bytes > MAX_ENCODED_BYTES
        ):
            break
        kept.append(part)
    dropped = range(len(kept), len(parts))
    numbers = tuple(part_nos[i] if part_nos is not None else i + 1 for i in dropped)
    return FittedParts(parts=tuple(kept), dropped_part_nos=numbers)
