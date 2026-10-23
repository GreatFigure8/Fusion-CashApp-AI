"""Spreadsheets: each non-empty XLSX sheet and each CSV becomes a TABLE part (CSV text)."""

import csv
import io
from datetime import date, datetime, time
from typing import Any, cast

from openpyxl import load_workbook

from cashapp.domain.artifacts import NormalisedPart, PartKind, UnsupportedReason
from cashapp.services.normalise.formats.text import decode_text
from cashapp.services.normalise.model import Item, Leaf, Reject, WalkContext, text_part
from cashapp.services.normalise.sniff import CSV


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    return str(value)


def _rows_to_csv(rows: list[list[str]]) -> str:
    while rows and not any(cell.strip() for cell in rows[-1]):
        rows.pop()
    trimmed: list[list[str]] = []
    for row in rows:
        end = len(row)
        while end and not row[end - 1].strip():
            end -= 1
        trimmed.append(row[:end])
    buffer = io.StringIO()
    csv.writer(buffer, lineterminator="\n").writerows(trimmed)
    return buffer.getvalue()


def handle_xlsx(item: Item, ctx: WalkContext) -> Leaf:
    workbook = load_workbook(io.BytesIO(item.data), read_only=True, data_only=True)
    parts: list[NormalisedPart] = []
    try:
        for sheet in workbook.worksheets:
            raw = cast(Any, sheet).iter_rows(values_only=True)
            rows = [[_cell(v) for v in cast(tuple[object, ...], row)] for row in raw]
            if not any(cell.strip() for row in rows for cell in row):
                continue
            parts.append(
                text_part(
                    PartKind.TABLE, CSV, _rows_to_csv(rows), f"{item.path}#sheet={sheet.title}"
                )
            )
    finally:
        workbook.close()
    if not parts:
        raise Reject(UnsupportedReason.EMPTY, "no non-empty sheet")
    return Leaf(parts)


def handle_csv(item: Item, ctx: WalkContext) -> Leaf:
    text = decode_text(item.data)
    if not text.strip():
        raise Reject(UnsupportedReason.EMPTY, "blank csv")
    return Leaf([text_part(PartKind.TABLE, CSV, text, item.path)])
